#!/usr/bin/env python3
"""
Generate one PMP 2026 content migration from the local question sheets.

Usage:
    python scripts/generate-pmp-content-migration.py <level>

Levels:
    case-studies    supabase/migrations/026_pmp_case_studies.sql
    full-exams      supabase/migrations/028_pmp_new_full_exams.sql
    agile           supabase/migrations/029_pmp_agile.sql
    tasks           supabase/migrations/030_pmp_task_exams.sql
    task-questions  supabase/migrations/031_pmp_task_new_questions.sql

Existing-question files (task exams) hold question_id | task_code. Every id must
be in an exam of SOURCE_EXAMS as 012 filled it; an id is linked, never copied.

Sheets live in src/data/exam/sheets/ (gitignored; never committed). Each
question sheet has an EN and an AR tab with the columns
    text | option_a .. option_e | correct_choice | explanation
read by header name, so tab order and extra columns do not matter. Any other
tab (e.g. "review") is ignored.

The run refuses on any EN/AR disagreement: row count, option count, or answer
key. New questions get database-assigned ids at migration time; the sheet
never supplies one.

Spec: docs/specs/spec-pmp-content-2026.md
"""

import importlib.util
import re
import string
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SHEETS = ROOT / "src" / "data" / "exam" / "sheets"
MIGRATIONS = ROOT / "supabase" / "migrations"

_spec = importlib.util.spec_from_file_location("content", Path(__file__).with_name("generate-content-migrations.py"))
content = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(content)

PMP_TRACK = content.PMP_TRACK
CONFIGS = {config["key"]: config for config in content.EXAM_CONFIGS}
OPTION_COLUMNS = [f"option_{letter}" for letter in "abcde"]
LETTERS = string.ascii_uppercase[: len(OPTION_COLUMNS)]

FULL_EXAMS = SHEETS / "full exams"
DOMAINS = SHEETS / "domain & tasks" / "PMP Domains"

# Reused task questions may only come from these exams, as 012 originally filled them:
# the 5 kept full exams, plus PMI Exam 1 (10) and PMI Exam 2 (14).
SOURCE_EXAMS = {17, 16, 15, 12, 13, 10, 14}
BANKS_MIGRATION = MIGRATIONS / "012_question_banks.sql"

# (tag, sheet folder, task code, English name, Arabic name), in display order.
TASKS = [
    ("people", "People", "I.1", "Develop a common vision", "تطوير رؤية مشتركة"),
    ("people", "People", "I.2", "Manage conflicts", "إدارة النزاعات"),
    ("people", "People", "I.3", "Lead the project team", "قيادة فريق المشروع"),
    ("people", "People", "I.4", "Engage stakeholders", "إشراك أصحاب المصلحة"),
    ("people", "People", "I.5", "Align stakeholder expectations", "مواءمة توقعات أصحاب المصلحة"),
    ("people", "People", "I.6", "Manage stakeholder expectations", "إدارة توقعات أصحاب المصلحة"),
    ("people", "People", "I.7", "Help ensure knowledge transfer", "المساعدة في ضمان نقل المعرفة"),
    ("people", "People", "I.8", "Plan and manage communication", "تخطيط الاتصالات وإدارتها"),
    ("process", "Process", "II.1", "Develop an integrated project management plan and plan delivery",
     "وضع خطة متكاملة لإدارة المشروع وتخطيط التسليم"),
    ("process", "Process", "II.2", "Develop and manage project scope", "تطوير نطاق المشروع وإدارته"),
    ("process", "Process", "II.3", "Help ensure value-based delivery", "المساعدة في ضمان التسليم القائم على القيمة"),
    ("process", "Process", "II.4", "Plan and manage resources", "تخطيط الموارد وإدارتها"),
    ("process", "Process", "II.5", "Plan and manage procurement", "تخطيط المشتريات وإدارتها"),
    ("process", "Process", "II.6", "Plan and manage finance", "تخطيط الشؤون المالية وإدارتها"),
    ("process", "Process", "II.7", "Plan and optimize quality of products/deliverables",
     "تخطيط جودة المنتجات/المخرجات وتحسينها"),
    ("process", "Process", "II.8", "Plan and manage schedule", "تخطيط الجدول الزمني وإدارته"),
    ("process", "Process", "II.9", "Evaluate project status", "تقييم حالة المشروع"),
    ("process", "Process", "II.10", "Manage project closure", "إدارة إغلاق المشروع"),
    ("environment", "Business Environment", "III.1", "Define and establish project governance",
     "تحديد حوكمة المشروع وإرساؤها"),
    ("environment", "Business Environment", "III.2", "Plan and manage project compliance", "تخطيط امتثال المشروع وإدارته"),
    ("environment", "Business Environment", "III.3", "Manage and control changes", "إدارة التغييرات وضبطها"),
    ("environment", "Business Environment", "III.4", "Remove impediments and manage issues", "إزالة العوائق وإدارة المشكلات"),
    ("environment", "Business Environment", "III.5", "Plan and manage risk", "تخطيط المخاطر وإدارتها"),
    ("environment", "Business Environment", "III.6", "Continuous improvement", "التحسين المستمر"),
    ("environment", "Business Environment", "III.7", "Support organizational change", "دعم التغيير المؤسسي"),
    ("environment", "Business Environment", "III.8", "Evaluate external business environment changes",
     "تقييم التغيرات في بيئة الأعمال الخارجية"),
]


def task_exam(tag, folder, code, name_en, name_ar):
    number = code.split(".")[1]
    return {"tag": tag, "display_order": int(number), "config": "domain", "task_code": code,
            "existing": (DOMAINS / folder, f"{code} *- existing.xlsx"),
            "new": (DOMAINS / folder, f"{code} *- new.xlsx"),
            "name_en": f"Task {number}: {name_en}", "name_ar": f"المهمة {number}: {name_ar}"}


LEVELS = {
    "case-studies": {
        "migration": "026_pmp_case_studies.sql",
        "title": "026 — PMP Case Studies exam.",
        # A case study is 5 consecutive questions; the first carries the scenario.
        "group_size": 5,
        "exams": [
            {"sheet": SHEETS / "domain & tasks" / "case_studies.xlsx", "tag": "casestudies", "display_order": 1,
             "config": "domain", "name_en": "Case Studies", "name_ar": "دراسات الحالة"},
        ],
    },
    "full-exams": {
        "migration": "028_pmp_new_full_exams.sql",
        "title": "028 — Three new PMP full exams, listed first.",
        "exams": [
            {"sheet": FULL_EXAMS / f"exam_{n}_pmp.xlsx", "tag": "full", "display_order": n, "config": "full",
             "name_en": f"PMP Core October 2026 ({n})", "name_ar": f"اختبار PMP Core أكتوبر 2026 ({n})"}
            for n in (1, 2, 3)
        ],
        # Kept full exams follow the new ones; archived PMP full exams move out of their way.
        "sql_after": [
            "UPDATE public.exams e SET display_order = e.display_order + 100",
            "FROM public.tags t",
            f"WHERE t.id = e.tag_id AND t.code = 'full' AND e.track_id = '{PMP_TRACK}' AND e.archived_at IS NOT NULL;",
            "",
            "UPDATE public.exams SET display_order = CASE id",
            "  WHEN 17 THEN 4  -- PMP Aug 2026",
            "  WHEN 16 THEN 5  -- PMP Core 2026",
            "  WHEN 15 THEN 6  -- PMI Exam 3",
            "  WHEN 12 THEN 7  -- PMP Full Exam ECO 1",
            "  WHEN 13 THEN 8  -- PMP Full Exam ECO 2",
            "END",
            "WHERE id IN (17, 16, 15, 12, 13);",
            "",
        ],
    },
    "agile": {
        "migration": "029_pmp_agile.sql",
        "title": "029 — Agile Methodology: a new Agile 1 exam, then the old Agile exam as Agile 2.",
        "exams": [
            {"sheet": SHEETS / "domain & tasks" / "agile_new.xlsx", "tag": "agile", "display_order": 1,
             "config": "domain", "name_en": "Agile 1", "name_ar": "أجايل 1"},
        ],
        "sql_after": [
            "-- The old Agile exam (200 questions) moves out of legacy, reopens as Agile 2.",
            "UPDATE public.exams",
            "SET tag_id        = (SELECT id FROM public.tags WHERE code = 'agile'),",
            "    display_order = 2,",
            "    name_en       = 'Agile 2',",
            "    name_ar       = 'أجايل 2',",
            "    archived_at   = NULL",
            "WHERE id = 41;",
            "",
        ],
    },
    "tasks": {
        "migration": "030_pmp_task_exams.sql",
        "title": "030 — 26 PMP task exams, holding their reused (existing) questions.",
        "exams": [task_exam(*task) for task in TASKS],
    },
    "task-questions": {
        "migration": "031_pmp_task_new_questions.sql",
        "title": "031 — New questions appended to the 26 PMP task exams, after their reused ones.",
        # Adds to exams 030 created instead of creating exams.
        "mode": "append",
        "exams": [task_exam(*task) for task in TASKS],
    },
}

# The break trigger rejects an exam with no question after a break index.
MIN_FULL_EXAM_QUESTIONS = 111


def cell(value):
    if value is None:
        return ""
    return str(value).replace("\r\n", "\n").strip()


def read_tab(workbook, name, path):
    rows = [row for row in workbook[name].iter_rows(values_only=True) if any(cell(c) for c in row)]
    header = [cell(c) for c in rows[0]]
    missing = {"text", "correct_choice", "explanation", *OPTION_COLUMNS[:2]} - set(header)
    if missing:
        raise SystemExit(f"{path.name} [{name}]: missing columns {sorted(missing)}")
    return [{h: cell(v) for h, v in zip(header, row) if h} for row in rows[1:]]


def parse_key(raw, where):
    letters = [part.strip().upper() for part in raw.split(",")]
    if not raw or any(len(letter) != 1 or letter not in LETTERS for letter in letters):
        raise SystemExit(f"{where}: correct_choice {raw!r} is not letters like 'B' or 'A,C'")
    return {LETTERS.index(letter) for letter in letters}


def options(row, where):
    values = [row.get(column, "") for column in OPTION_COLUMNS]
    count = len(values) - next((i for i, v in enumerate(reversed(values)) if v), len(values))
    if not all(values[:count]) or count < 2:
        raise SystemExit(f"{where}: options must be at least 2 and contiguous from option_a")
    return values[:count]


def read_questions(path):
    """Paired EN/AR questions from one sheet, in row order."""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    en, ar = read_tab(workbook, "EN", path), read_tab(workbook, "AR", path)
    if len(en) != len(ar):
        raise SystemExit(f"{path.name}: EN has {len(en)} rows, AR has {len(ar)}")

    questions = []
    for number, (row_en, row_ar) in enumerate(zip(en, ar), start=2):
        where = f"{path.name} row {number}"
        for field in ("text", "explanation"):
            if not row_en[field] or not row_ar[field]:
                raise SystemExit(f"{where}: empty {field}")
        options_en, options_ar = options(row_en, f"{where} EN"), options(row_ar, f"{where} AR")
        if len(options_en) != len(options_ar):
            raise SystemExit(f"{where}: EN has {len(options_en)} options, AR has {len(options_ar)}")
        key_en, key_ar = parse_key(row_en["correct_choice"], f"{where} EN"), parse_key(row_ar["correct_choice"], f"{where} AR")
        if key_en != key_ar:
            raise SystemExit(f"{where}: answer key differs, EN {row_en['correct_choice']} vs AR {row_ar['correct_choice']}")
        if max(key_en) >= len(options_en):
            raise SystemExit(f"{where}: correct_choice points past the last option")
        questions.append({
            "text_en": row_en["text"], "text_ar": row_ar["text"],
            "explanation_en": row_en["explanation"], "explanation_ar": row_ar["explanation"],
            "choices": [(text_ar, text_en, position in key_en)
                        for position, (text_en, text_ar) in enumerate(zip(options_en, options_ar))],
        })
    return questions


def find_sheet(folder, pattern):
    matches = sorted(folder.glob(pattern))
    if len(matches) != 1:
        raise SystemExit(f"expected one sheet matching {pattern!r} in {folder}, found {len(matches)}")
    return matches[0]


def read_existing(path, task_code):
    """Question ids of an existing-question file, in row order."""
    rows = [r for r in openpyxl.load_workbook(path, read_only=True).worksheets[0].iter_rows(values_only=True)
            if any(cell(c) for c in r)]
    if [cell(c) for c in rows[0][:2]] != ["question_id", "task_code"]:
        raise SystemExit(f"{path.name}: expected columns question_id | task_code")
    ids = []
    for number, row in enumerate(rows[1:], start=2):
        if cell(row[1]) != task_code:
            raise SystemExit(f"{path.name} row {number}: task_code {cell(row[1])!r}, expected {task_code!r}")
        if not cell(row[0]).isdigit():
            raise SystemExit(f"{path.name} row {number}: question_id {cell(row[0])!r} is not a number")
        ids.append(int(cell(row[0])))
    if len(set(ids)) != len(ids):
        raise SystemExit(f"{path.name}: duplicate question ids")
    return ids


def source_question_ids():
    """Question ids of SOURCE_EXAMS as 012 inserted them (before 027 trimmed any)."""
    rows = re.findall(r"^  \((\d+), \d+, (\d+)\)[,;]$", BANKS_MIGRATION.read_text(encoding="utf-8"), re.MULTILINE)
    ids = {int(question) for exam, question in rows if int(exam) in SOURCE_EXAMS}
    if not ids:
        raise SystemExit(f"no exam_questions rows found for {sorted(SOURCE_EXAMS)} in {BANKS_MIGRATION.name}")
    return ids


def render_level(level, exams):
    """One migration inserting every exam of a level: its new questions, choices and links.
    An int in an exam's question list is an existing question id, linked rather than copied."""
    s = content.sql_string
    append = level.get("mode") == "append"
    exam_rows, question_rows, choice_rows, link_rows = [], [], [], []
    for exam_no, (exam, questions) in enumerate(exams):
        config = CONFIGS[exam["config"]]
        base = exam.get("base_count", 0)
        description_ar, description_en = content.describe(exam["config"], base + len(questions), config)
        if append:
            exam_rows.append([exam_no, s(exam["tag"]), exam["display_order"], base, base + len(questions),
                              s(description_ar), s(description_en)])
        else:
            exam_rows.append([exam_no, s(exam["tag"]), exam["display_order"], config["id"], s(exam["name_ar"]),
                              s(exam["name_en"]), s(description_ar), s(description_en), len(questions)])
        for index, q in enumerate(questions):
            if isinstance(q, int):
                link_rows.append([exam_no, index, q])
                continue
            question_rows.append([exam_no, index, s(q["text_ar"]), s(q["text_en"]), s(q["explanation_ar"]),
                                  s(q["explanation_en"]), sum(1 for c in q["choices"] if c[2])])
            for position, (text_ar, text_en, correct) in enumerate(q["choices"]):
                choice_rows.append([exam_no, index, position, s(text_ar), s(text_en), "true" if correct else "false"])

    def temp_inserts(table, columns, rows):
        return content.render_inserts(table, columns, rows, lambda row: ", ".join(map(str, row)), schema="")

    lines = [
        "-- " + "=" * 77,
        f"-- {level['title']}",
        "--",
        "-- GENERATED by scripts/generate-pmp-content-migration.py - do not edit by hand.",
        f"-- {len(exams)} exam(s), {len(question_rows)} new questions. Question ids come from the",
        "-- questions sequence, consecutive in exam then sheet order.",
        *([f"-- {len(link_rows)} existing questions are linked, not copied."] if link_rows else []),
        "-- " + "=" * 77,
        "",
        "BEGIN;",
        "",
    ]
    if append:
        lines += [
            "CREATE TEMP TABLE target_exams (",
            "  exam_no         smallint PRIMARY KEY,",
            "  tag_code        text     NOT NULL,",
            "  display_order   smallint NOT NULL,",
            "  base_count      smallint NOT NULL,",
            "  question_count  smallint NOT NULL,",
            "  description_ar  text     NOT NULL,",
            "  description_en  text     NOT NULL,",
            "  id              smallint",
            ") ON COMMIT DROP;",
            "",
        ]
    else:
        lines += [
        "CREATE TEMP TABLE new_exams (",
        "  exam_no         smallint PRIMARY KEY,",
        "  tag_code        text     NOT NULL,",
        "  display_order   smallint NOT NULL,",
        "  config_id       smallint NOT NULL,",
        "  name_ar         text     NOT NULL,",
        "  name_en         text     NOT NULL UNIQUE,",
        "  description_ar  text     NOT NULL,",
        "  description_en  text     NOT NULL,",
        "  question_count  smallint NOT NULL,",
        "  id              smallint",
        ") ON COMMIT DROP;",
        "",
    ]
    if question_rows:
        lines += [
        "CREATE TEMP TABLE new_questions (",
        "  exam_no         smallint NOT NULL,",
        "  question_index  smallint NOT NULL,",
        "  text_ar         text     NOT NULL,",
        "  text_en         text     NOT NULL,",
        "  explanation_ar  text     NOT NULL,",
        "  explanation_en  text     NOT NULL,",
        "  answer_count    smallint NOT NULL,",
        "  id              integer,",
        "  PRIMARY KEY (exam_no, question_index)",
        ") ON COMMIT DROP;",
        "",
        "CREATE TEMP TABLE new_choices (",
        "  exam_no         smallint NOT NULL,",
        "  question_index  smallint NOT NULL,",
        "  position        smallint NOT NULL,",
        "  text_ar         text     NOT NULL,",
        "  text_en         text     NOT NULL,",
        "  is_correct      boolean  NOT NULL",
        ") ON COMMIT DROP;",
        "",
        ]
    if link_rows:
        lines += [
            "CREATE TEMP TABLE new_links (",
            "  exam_no         smallint NOT NULL,",
            "  question_index  smallint NOT NULL,",
            "  question_id     integer  NOT NULL",
            ") ON COMMIT DROP;",
            "",
        ]
    if append:
        lines += temp_inserts("target_exams", ["exam_no", "tag_code", "display_order", "base_count", "question_count",
                                               "description_ar", "description_en"], exam_rows)
    else:
        lines += temp_inserts("new_exams", ["exam_no", "tag_code", "display_order", "config_id", "name_ar", "name_en",
                                            "description_ar", "description_en", "question_count"], exam_rows)
    lines += temp_inserts("new_questions", ["exam_no", "question_index", "text_ar", "text_en", "explanation_ar",
                                            "explanation_en", "answer_count"], question_rows)
    lines += temp_inserts("new_choices", ["exam_no", "question_index", "position", "text_ar", "text_en", "is_correct"],
                          choice_rows)
    lines += temp_inserts("new_links", ["exam_no", "question_index", "question_id"], link_rows)
    if question_rows:
        lines += [
        "-- Consecutive ids from the questions sequence, in exam then sheet order.",
        "WITH base AS (SELECT nextval(pg_get_serial_sequence('public.questions', 'id')) AS id),",
        "     ordered AS (SELECT exam_no, question_index,",
        "                        row_number() OVER (ORDER BY exam_no, question_index) - 1 AS n",
        "                 FROM new_questions)",
        "UPDATE new_questions q SET id = base.id + o.n",
        "FROM base, ordered o",
        "WHERE o.exam_no = q.exam_no AND o.question_index = q.question_index;",
        "SELECT setval(pg_get_serial_sequence('public.questions', 'id'), (SELECT max(id) FROM new_questions));",
        "",
        "INSERT INTO public.questions (id, type, text_ar, text_en, explanation_ar, explanation_en, answer_count)",
        "SELECT id, 'multiple-choice', text_ar, text_en, explanation_ar, explanation_en, answer_count",
        "FROM new_questions ORDER BY id;",
        "",
        "INSERT INTO public.choices (question_id, position, text_ar, text_en, is_correct)",
        "SELECT q.id, c.position, c.text_ar, c.text_en, c.is_correct",
        "FROM new_choices c JOIN new_questions q USING (exam_no, question_index);",
        "",
        ]
    if append:
        lines += [
            "-- Each target is the active PMP exam at (tag, display_order), still holding exactly",
            "-- base_count questions; anything else means the inputs drifted, so stop.",
            "UPDATE target_exams x SET id = e.id",
            "FROM public.exams e JOIN public.tags t ON t.id = e.tag_id",
            f"WHERE t.code = x.tag_code AND e.display_order = x.display_order AND e.track_id = {s(PMP_TRACK)}",
            "  AND e.archived_at IS NULL;",
            "",
            "DO $$",
            "DECLARE",
            "  v_problem TEXT;",
            "BEGIN",
            "  SELECT string_agg(format('%s #%s', x.tag_code, x.display_order), ', ') INTO v_problem",
            "  FROM target_exams x",
            "  WHERE x.id IS NULL",
            "     OR (SELECT count(*) FROM public.exam_questions eq WHERE eq.exam_id = x.id) <> x.base_count;",
            "  IF v_problem IS NOT NULL THEN",
            "    RAISE EXCEPTION 'Target exam missing or not at its expected question count: %', v_problem;",
            "  END IF;",
            "END",
            "$$;",
            "",
            "INSERT INTO public.exam_questions (exam_id, question_index, question_id)",
            "SELECT x.id, x.base_count + q.question_index, q.id FROM new_questions q JOIN target_exams x USING (exam_no);",
            "",
            "UPDATE public.exams e",
            "SET question_count = x.question_count, description_ar = x.description_ar, description_en = x.description_en",
            "FROM target_exams x",
            "WHERE e.id = x.id;",
            "",
        ]
        lines += level.get("sql_after", [])
        lines += ["COMMIT;", ""]
        return "\n".join(lines)

    lines += [
        "-- An unknown tag code yields a NULL tag_id, which the NOT NULL column rejects.",
        "WITH inserted AS (",
        "  INSERT INTO public.exams (track_id, config_id, tag_id, display_order, name_ar, name_en,",
        "                            description_ar, description_en, question_count)",
        f"  SELECT {s(PMP_TRACK)}, config_id, (SELECT id FROM public.tags WHERE code = tag_code), display_order,",
        "         name_ar, name_en, description_ar, description_en, question_count",
        "  FROM new_exams ORDER BY exam_no",
        "  RETURNING id, name_en",
        ")",
        "UPDATE new_exams n SET id = i.id FROM inserted i WHERE i.name_en = n.name_en;",
        "",
    ]
    if question_rows:
        lines += [
            "INSERT INTO public.exam_questions (exam_id, question_index, question_id)",
            "SELECT n.id, q.question_index, q.id FROM new_questions q JOIN new_exams n USING (exam_no);",
            "",
        ]
    if link_rows:
        lines += [
            "INSERT INTO public.exam_questions (exam_id, question_index, question_id)",
            "SELECT n.id, l.question_index, l.question_id FROM new_links l JOIN new_exams n USING (exam_no);",
            "",
        ]
    lines += level.get("sql_after", [])
    lines += ["COMMIT;", ""]
    return "\n".join(lines)


def main():
    if len(sys.argv) != 2 or sys.argv[1] not in LEVELS:
        raise SystemExit(f"usage: generate-pmp-content-migration.py <{'|'.join(LEVELS)}>")
    level = LEVELS[sys.argv[1]]

    exams = []
    sources = None
    used_by = {}
    for exam in level["exams"]:
        if level.get("mode") == "append":
            path = find_sheet(*exam["new"])
            base = len(read_existing(find_sheet(*exam["existing"]), exam["task_code"]))
            exams.append(({**exam, "base_count": base}, read_questions(path)))
            continue
        if "existing" in exam:
            path = find_sheet(*exam["existing"])
            ids = read_existing(path, exam["task_code"])
            sources = sources or source_question_ids()
            outside = [question_id for question_id in ids if question_id not in sources]
            if outside:
                raise SystemExit(f"{path.name}: ids not in exams {sorted(SOURCE_EXAMS)}: {outside}")
            for question_id in ids:
                used_by.setdefault(question_id, []).append(exam["task_code"])
            exams.append((exam, ids))
            continue
        if not exam["sheet"].exists():
            raise SystemExit(f"missing sheet: {exam['sheet']}")
        questions = read_questions(exam["sheet"])
        group = level.get("group_size")
        if group and len(questions) % group:
            raise SystemExit(f"{exam['sheet'].name}: {len(questions)} questions is not a whole number of {group}-question case studies")
        if exam["config"] == "full" and len(questions) < MIN_FULL_EXAM_QUESTIONS:
            raise SystemExit(f"{exam['sheet'].name}: a full exam needs at least {MIN_FULL_EXAM_QUESTIONS} questions")
        exams.append((exam, questions))

    for question_id, codes in sorted(used_by.items()):
        if len(codes) > 1:
            print(f"warning: question {question_id} is reused by tasks {', '.join(codes)}")

    path = MIGRATIONS / level["migration"]
    content.write(path, render_level(level, exams))
    for exam, questions in exams:
        new = [q for q in questions if not isinstance(q, int)]
        choices = sum(len(q["choices"]) for q in new)
        linked = len(questions) - len(new)
        print(f"{path.name}: {exam['name_en']}: {len(new)} new questions, {choices} choices"
              + (f", {linked} linked" if linked else "")
              + (f", after {exam['base_count']} existing" if "base_count" in exam else ""))
    return 0


if __name__ == "__main__":
    sys.exit(main())
